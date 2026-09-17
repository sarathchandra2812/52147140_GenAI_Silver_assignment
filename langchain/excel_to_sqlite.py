from __future__ import annotations

import logging
import os
import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s - %(funcName)s:%(lineno)d - %(message)s",
)
logger = logging.getLogger("excel_to_sqlite")

# Default paths you can edit directly.
BASE_DIR = Path(__file__).resolve().parent
INPUT_XLSX = BASE_DIR / "smb_intelligence_real_world_sweden.xlsx"
OUTPUT_DB = BASE_DIR / "sweden_data.db"

# Optional indexes for better query performance.
INDEX_PLAN: dict[str, list[str]] = {
    "competitor_pricing": ["date", "competitor_name", "product_name"],
    "sales_data": ["date", "retailer", "product_name"],
    "inventory_status": ["snapshot_date", "retailer", "product_name"],
    "customer_feedback": ["date", "retailer", "product_name"],
}


def _truthy(name: str, default: bool) -> bool:
    raw = os.getenv(name)
    if raw is None:
        logger.info("Env %s not set, using default=%s", name, default)
        return default
    resolved = raw.strip().lower() in {"1", "true", "yes", "y", "on"}
    logger.info("Env %s=%r resolved to %s", name, raw, resolved)
    return resolved


def _resolve_path(name: str, default_path: Path) -> Path:
    override = os.getenv(name)
    if not override:
        logger.info("Env %s not set, using default path: %s", name, default_path)
        return default_path
    resolved_path = Path(override).expanduser().resolve()
    logger.info("Env %s override path resolved to: %s", name, resolved_path)
    return resolved_path


def _quote_ident(identifier: str) -> str:
    return '"' + identifier.replace('"', '""') + '"'


def _drop_existing_user_tables(conn: sqlite3.Connection) -> None:
    rows = conn.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'").fetchall()
    logger.info("Found %d existing user table(s) to drop", len(rows))
    for (table_name,) in rows:
        logger.info("Dropping table: %s", table_name)
        conn.execute(f"DROP TABLE IF EXISTS {_quote_ident(table_name)}")


def _write_readme_table(conn: sqlite3.Connection, workbook_path: Path) -> None:
    logger.info("Loading workbook to extract README sheet: %s", workbook_path)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "README" not in workbook.sheetnames:
        logger.warning("README sheet not found in workbook")
        return

    readme_rows: list[dict[str, str]] = []
    readme_sheet = workbook["README"]
    for row in readme_sheet.iter_rows(values_only=True):
        if not row or row[0] is None:
            continue
        item = str(row[0]).strip()
        detail = "" if len(row) < 2 or row[1] is None else str(row[1]).strip()
        readme_rows.append({"item": item, "detail": detail})

    if readme_rows:
        logger.info("Writing README table with %d row(s)", len(readme_rows))
        pd.DataFrame(readme_rows).to_sql("README", conn, if_exists="replace", index=False)
    else:
        logger.warning("README sheet exists but no non-empty rows were found")


def _import_sheets(conn: sqlite3.Connection, workbook_path: Path, include_readme: bool) -> None:
    logger.info("Reading workbook sheets from: %s", workbook_path)
    sheet_names = pd.ExcelFile(workbook_path).sheet_names
    logger.info("Discovered sheets: %s", ", ".join(sheet_names))
    for sheet_name in sheet_names:
        if sheet_name == "README" and include_readme:
            logger.info("Skipping README sheet import because include_readme=%s", include_readme)
            continue
        dataframe = pd.read_excel(workbook_path, sheet_name=sheet_name)
        logger.info("Importing sheet %s with %d row(s), %d column(s)", sheet_name, len(dataframe), len(dataframe.columns))
        dataframe.to_sql(sheet_name, conn, if_exists="replace", index=False)


def _create_indexes(conn: sqlite3.Connection) -> None:
    existing_tables = {row[0] for row in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    logger.info("Creating indexes for available tables: %s", ", ".join(sorted(existing_tables)))

    for table_name, columns in INDEX_PLAN.items():
        if table_name not in existing_tables:
            logger.warning("Skipping index plan for missing table: %s", table_name)
            continue

        table_info = conn.execute(f"PRAGMA table_info({_quote_ident(table_name)})").fetchall()
        existing_columns = {row[1] for row in table_info}

        for column_name in columns:
            if column_name not in existing_columns:
                logger.warning("Skipping index for %s.%s (column not found)", table_name, column_name)
                continue
            index_name = f"idx_{table_name}_{column_name}"
            logger.info("Ensuring index exists: %s on %s(%s)", index_name, table_name, column_name)
            conn.execute(f"CREATE INDEX IF NOT EXISTS {_quote_ident(index_name)} " f"ON {_quote_ident(table_name)}({_quote_ident(column_name)})")


def main() -> None:
    logger.info("Starting Excel to SQLite import workflow")
    input_xlsx = _resolve_path("SWEDEN_XLSX_PATH", INPUT_XLSX)
    output_db = _resolve_path("SWEDEN_DB_PATH", OUTPUT_DB)
    include_readme = _truthy("INCLUDE_README", True)
    drop_existing = _truthy("DROP_EXISTING_TABLES", True)

    logger.info("Input workbook: %s", input_xlsx)
    logger.info("Output database: %s", output_db)
    logger.info("Options -> include_readme=%s, drop_existing_tables=%s", include_readme, drop_existing)

    if not input_xlsx.exists():
        logger.error("Input workbook not found: %s", input_xlsx)
        raise FileNotFoundError(f"Input workbook not found: {input_xlsx}")

    output_db.parent.mkdir(parents=True, exist_ok=True)
    logger.info("Ensured output directory exists: %s", output_db.parent)

    with sqlite3.connect(output_db) as conn:
        logger.info("Connected to SQLite database: %s", output_db)
        if drop_existing:
            _drop_existing_user_tables(conn)
        else:
            logger.info("Skipping drop of existing tables")

        if include_readme:
            _write_readme_table(conn, input_xlsx)
        else:
            logger.info("Skipping README table import")

        _import_sheets(conn, input_xlsx, include_readme=include_readme)
        _create_indexes(conn)

        tables = [row[0] for row in conn.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name")]
        logger.info("Import completed. Final table count: %d", len(tables))
        print("Created/updated database:", output_db)
        print("Tables:", ", ".join(tables))
        for table_name in tables:
            count = conn.execute(f"SELECT COUNT(*) FROM {_quote_ident(table_name)}").fetchone()[0]
            logger.info("Table %s row count: %d", table_name, count)
            print(f"  {table_name}: {count}")

    logger.info("Excel to SQLite workflow finished successfully")


if __name__ == "__main__":
    main()
